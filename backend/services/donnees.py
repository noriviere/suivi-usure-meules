import os
import statistics
from collections import Counter
from openpyxl import load_workbook

from .. import config
FEUILLE_DATA = "Data"
COLONNE_CLE = "n° article meule"       # une ligne de données a toujours un article

_cache = {"mtime": None, "lignes": []}


def norm(texte):
    """Ignore majuscules, espaces multiples et retours à la ligne."""
    return " ".join(str(texte).split()).lower() if texte is not None else ""


def lire_data():
    """Lignes de la feuille Data (dict par en-tête normalisé).
    Relu seulement si le fichier a changé depuis la dernière lecture."""
    mtime = os.path.getmtime(config.FICHIER)
    if _cache["mtime"] == mtime:
        return _cache["lignes"]
    try:
        wb = load_workbook(config.FICHIER, read_only=True, data_only=True)
    except PermissionError:
        return _cache["lignes"]            # fichier verrouillé : on garde l'ancien état
    ws = wb[FEUILLE_DATA]
    rows = ws.iter_rows(values_only=True)
    entetes = [norm(entete) for entete in next(rows)]
    # On repère les lignes remplies grâce à l'article (et non à la 1re colonne,
    # qui peut changer quand on réorganise le tableau Excel).
    cle = entetes.index(COLONNE_CLE) if COLONNE_CLE in entetes else 0
    lignes = [dict(zip(entetes, ligne_brute)) for ligne_brute in rows
              if len(ligne_brute) > cle and ligne_brute[cle] not in (None, "")]
    wb.close()
    _cache["mtime"], _cache["lignes"] = mtime, lignes
    return lignes


def matricules():
    """Matricules déjà utilisés dans le fichier (monteurs et démonteurs),
    du plus fréquent au plus rare. Les numéros sont complétés à 4 chiffres
    ("534" -> "0534")."""
    compte = Counter()
    for ligne in lire_data():
        for cle, valeur in ligne.items():
            if not cle.startswith("matricule") or valeur in (None, ""):
                continue
            if isinstance(valeur, float) and valeur.is_integer():
                valeur = int(valeur)
            texte = str(valeur).strip()
            if texte.isdigit():
                texte = texte.zfill(4)
            compte[texte] += 1
    return [matricule for matricule, _ in compte.most_common()]


def mediane(article, machine, matiere, minimum=5):
    """Médiane de 'Nb faces/meule de 89 mm' pour la face A et la face B.
    On part du critère le plus précis et on élargit tant qu'il y a moins
    de `minimum` cas. Les lignes X sont incluses."""
    lignes = lire_data()
    niveaux = [
        ("article + machine + matière", {"n° article meule": article, "n° de machine": machine, "matière": matiere}),
        ("article + matière",           {"n° article meule": article, "matière": matiere}),
        ("article seul",                {"n° article meule": article}),
    ]
    resultat = {}
    for face in ("A", "B"):
        resultat[face] = None
        repli = None
        for nom, criteres in niveaux:
            valeurs = [
                ligne["nb faces/meule de 89 mm"] for ligne in lignes
                if norm(ligne.get("face")) == norm(face)
                and all(norm(ligne.get(critere)) == norm(valeur) for critere, valeur in criteres.items())
                and isinstance(ligne.get("nb faces/meule de 89 mm"), (int, float))
            ]
            if not valeurs:
                continue
            info = {"mediane": round(statistics.median(valeurs)), "nb_cas": len(valeurs),
                    "critere": nom, "fiable": len(valeurs) >= minimum}
            if info["fiable"]:
                resultat[face] = info
                break
            repli = repli or info
        else:
            resultat[face] = repli
    return resultat


_cache_t = {"mtime": None, "donnees": None}


def _json(valeur):
    if hasattr(valeur, "isoformat"):
        return valeur.isoformat()[:10]
    if isinstance(valeur, bool) or valeur is None:
        return valeur
    if isinstance(valeur, (int, float)):
        return valeur
    return str(valeur)


def tableau():
    """Tout le contenu de la feuille Data pour la page de consultation :
    {"entetes": [...], "cles": [...], "lignes": [[n° de ligne Excel, valeur, ...], ...]}.
    Mis en cache tant que le fichier ne change pas."""
    mtime = os.path.getmtime(config.FICHIER)
    if _cache_t["mtime"] == mtime and _cache_t["donnees"]:
        return _cache_t["donnees"]
    try:
        wb = load_workbook(config.FICHIER, read_only=True, data_only=True)
    except PermissionError:
        return _cache_t["donnees"] or {"entetes": [], "cles": [], "lignes": []}
    rows = wb[FEUILLE_DATA].iter_rows(values_only=True)
    brut = next(rows)
    nb = max((position for position, entete in enumerate(brut) if entete not in (None, "")), default=-1) + 1
    entetes = [" ".join(str(entete).split()) for entete in brut[:nb]]
    cles = [norm(entete) for entete in entetes]
    cle = cles.index(COLONNE_CLE) if COLONNE_CLE in cles else 0
    lignes = []
    for numero_ligne, ligne_brute in enumerate(rows, start=2):
        if len(ligne_brute) <= cle or ligne_brute[cle] in (None, ""):
            continue
        lignes.append([numero_ligne] + [_json(valeur) for valeur in list(ligne_brute[:nb]) + [None] * (nb - len(ligne_brute))])
    wb.close()
    resultat = {"entetes": entetes, "cles": cles, "lignes": lignes}
    _cache_t["mtime"], _cache_t["donnees"] = mtime, resultat
    return resultat
