"""Écriture en mode DÉMO : ajoute la ligne dans le classeur de fausses données avec openpyxl (pas d'Excel, pas de COM).
Les colonnes que le vrai classeur calcule par formules (usure, faces par mm...) sont calculées ici en Python.
Ce module n'est utilisé que si MEULES_DEMO=1 ; en production l'écriture passe toujours par ecriture.py (Excel)."""
import time

from openpyxl import load_workbook

from .. import config
from .donnees import norm
from .ecriture import FichierOccupe, VERROU, _Etape
from . import ecriture as _ecriture

PAUSE_ETAPE_S = 0.35        # petite pause par étape pour que la barre de chargement soit visible en démo
ETAPES = ["lancement_excel", "ouverture", "ajout_ligne", "formules", "formats", "valeurs", "calcul", "sauvegarde", "fermeture"]


def _calculer(ligne, meules, codes_lame):
    """Équivalent Python des formules du classeur (clés = noms de colonnes normalisés)."""
    sortie = {}
    article = str(ligne.get("n° article meule") or "")
    meule = meules.get(article, {})
    sortie["fournisseur"], sortie["ø meule"] = meule.get("fournisseur"), meule.get("diametre")
    sortie["code lame"] = codes_lame.get(norm(str(ligne.get("format") or "") + str(ligne.get("matière") or "")))
    debut, fin = ligne.get("compteur début"), ligne.get("compteur fin")
    haut_i, haut_f = ligne.get("hauteur meule initiale, mm"), ligne.get("hauteur meule finale, mm")
    brut = (fin - debut) if isinstance(fin, (int, float)) and isinstance(debut, (int, float)) and fin != 0 else None
    usure = (haut_i - haut_f) if isinstance(haut_i, (int, float)) and isinstance(haut_f, (int, float)) and haut_i != 0 else None
    sortie["nombre de face brut"] = brut
    sortie["ht.usée, mm"] = usure
    sortie["meule démontée avant usure totale"] = "X" if isinstance(haut_f, (int, float)) and haut_f > 2 else None
    faces_mm = brut / usure if brut is not None and usure else None
    sortie["nb faces/mm"] = faces_mm
    sortie["nb faces/meule de 89 mm"] = faces_mm * 89 if faces_mm is not None else None
    date = ligne.get("date montage")
    sortie["mois"], sortie["année"] = (date.month, date.year) if hasattr(date, "month") else (None, None)
    sortie["ht usée par face, mm"] = usure / brut if brut and usure is not None else None
    return sortie


def _ajouter(feuille, valeurs, meules, codes_lame, tolerant=False):
    """Ajoute une ligne à la fin de la feuille et agrandit son tableau Excel. Renvoie (n° de ligne, calculés, ignorés)."""
    entetes = [norm(c.value) for c in feuille[1]]
    brut = {norm(nom): valeur for nom, valeur in valeurs.items()}
    ignores = [nom for nom in valeurs if norm(nom) not in entetes and tolerant]
    ligne = dict(brut)
    ligne.update({k: v for k, v in _calculer(ligne, meules, codes_lame).items() if k in entetes})
    numero = feuille.max_row + 1
    for colonne, entete in enumerate(entetes, start=1):
        if entete in ligne and ligne[entete] is not None:
            feuille.cell(row=numero, column=colonne, value=ligne[entete])
    for table in feuille.tables.values():
        debut = table.ref.split(":")[0]
        fin_colonne = "".join(c for c in table.ref.split(":")[1] if c.isalpha())
        table.ref = f"{debut}:{fin_colonne}{numero}"
    calcules = {}
    for entete in entetes:
        if entete not in brut and entete in ligne:
            calcules[entete] = "" if ligne[entete] is None else ligne[entete]
    return numero, calcules, ignores


def ajouter_ligne(valeurs, defaut=False, progression=None, mesures=None):
    """Même contrat que ecriture.ajouter_ligne : renvoie (n° de ligne, colonnes calculées, infos onglet défaut)."""
    cfg = _ecriture.CONFIG.get("defaut") or {}
    if progression:
        progression("attente")
    with VERROU:
        _ecriture.MESURES.clear()
        _ecriture._PROGRESSION = progression
        try:
            for etape in ETAPES[:3]:
                with _Etape(etape):
                    time.sleep(PAUSE_ETAPE_S)
            try:
                wb = load_workbook(config.FICHIER)
            except PermissionError:
                raise FichierOccupe("Le classeur de démonstration est ouvert ailleurs.")
            meules = {}
            for ligne in wb["Liste Meule"].iter_rows(min_row=2, values_only=True):
                if ligne[0]:
                    meules[str(ligne[0])] = {"fournisseur": ligne[3], "diametre": ligne[4]}
            codes_lame = {}
            for ligne in wb["Formats & matière"].iter_rows(min_row=2, values_only=True):
                if ligne[2] is not None:
                    codes_lame[norm(str(ligne[2]) + str(ligne[3] or ""))] = ligne[1]
            for etape in ETAPES[3:6]:
                with _Etape(etape):
                    time.sleep(PAUSE_ETAPE_S)
            numero, calcules, info_defaut = None, {}, None
            if not defaut or cfg.get("garder_dans_data", True):
                numero, calcules, _ = _ajouter(wb["Data"], valeurs, meules, codes_lame)
            if defaut:
                feuille = next((ws for ws in wb.worksheets if norm(ws.title).replace("é", "e") == norm(cfg["onglet"]).replace("é", "e")), None)
                if feuille is None:
                    raise ValueError(f"Onglet « {cfg['onglet']} » introuvable dans le classeur.")
                ligne2, _, ignores = _ajouter(feuille, valeurs, meules, codes_lame, tolerant=True)
                info_defaut = {"onglet": feuille.title, "ligne": ligne2, "colonnes_absentes": ignores}
                numero = numero or ligne2
            with _Etape("calcul"):
                time.sleep(PAUSE_ETAPE_S)
            with _Etape("sauvegarde"):
                wb.save(config.FICHIER)
            with _Etape("fermeture"):
                wb.close()
                time.sleep(PAUSE_ETAPE_S)
            return numero, calcules, info_defaut
        finally:
            _ecriture._PROGRESSION = None
            if mesures is not None:
                mesures.update(_ecriture.MESURES)
