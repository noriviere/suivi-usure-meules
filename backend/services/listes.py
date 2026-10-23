import os
import re
from openpyxl import load_workbook

from .. import config
from .donnees import norm, matricules

_cache = {"mtime": None, "formats": [], "meules": []}


def _txt(valeur):
    if valeur is None:
        return ""
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur))               # 2.0 -> "2"
    return str(valeur).strip()


def _liste(texte):
    """'s1;s2;s7' -> ['s1', 's2', 's7'] (accepte aussi / et ,)"""
    return [element.strip() for element in re.split(r"[;/,]", texte) if element.strip()]


def _lire(wb, feuille):
    rows = wb[feuille].iter_rows(values_only=True)
    entetes = [norm(entete) for entete in next(rows)]
    return [dict(zip(entetes, ligne_brute)) for ligne_brute in rows if ligne_brute and ligne_brute[0] is not None]


def charger():
    mtime = os.path.getmtime(config.FICHIER)
    if _cache["mtime"] == mtime:
        return _cache
    try:
        wb = load_workbook(config.FICHIER, read_only=True, data_only=True)
    except PermissionError:
        return _cache                    # fichier verrouillé : on garde l'ancien état
    formats = []
    for ligne in _lire(wb, "Formats & matière"):
        fmt, mat = _txt(ligne.get("format")), _txt(ligne.get("matière"))
        meules = _liste(_txt(ligne.get("meule")))
        formats.append({
            "cle": norm(fmt + mat),      # même clé que la formule Excel Format&Matière
            "format": fmt,
            "matiere": mat,
            "code": _txt(ligne.get("code article")),
            "ma5": _txt(ligne.get("ma5")),
            "meule": meules[0] if meules else "",
            "meules": meules,            # meule(s) prévue(s) pour ce format + matière
            "machines": _liste(_txt(ligne.get("machine"))),
        })
    meules = []
    for ligne in _lire(wb, "Liste Meule"):
        code = _txt(ligne.get("code article"))
        if code:
            meules.append({
                "code": code,
                "etat": _txt(ligne.get("etat")),
                "fournisseur": _txt(ligne.get("fournisseur")),
                "diametre": _txt(ligne.get("diamètre meule")),
            })
    wb.close()
    _cache.update(mtime=mtime, formats=formats, meules=meules)
    return _cache


def ligne_format(fmt, matiere):
    cle = norm(_txt(fmt) + _txt(matiere))
    for ref_format in charger()["formats"]:
        if ref_format["cle"] == cle:
            return ref_format
    return None


def listes_deroulantes():
    referentiel = charger()
    return {
        "N° article meule": [meule["code"] for meule in referentiel["meules"]],
        "Format": list(dict.fromkeys(ref_format["format"] for ref_format in referentiel["formats"])),
        "Matière": list(dict.fromkeys(ref_format["matiere"] for ref_format in referentiel["formats"])),
        "N° de Machine": sorted({machine for ref_format in referentiel["formats"] for machine in ref_format["machines"]}, key=norm),
        "matricules": matricules(),
    }
