"""Adapter openpyxl du module d'écriture : sans Excel ni COM (Linux, macOS, Codespaces).
Sert au mode démo (MEULES_DEMO=1, avec une petite pause par étape) et aux tests (sans pause).
Les colonnes que le vrai classeur calcule par formules (usure, faces par mm...) sont calculées ici en Python ;
les formules du classeur restent la référence. En production, l'écriture passe toujours par classeur_excel.py."""
import time

from openpyxl import load_workbook

from .donnees import norm
from .ecriture import FEUILLE, TABLEAU, FichierOccupe


def _calculer(ligne, meules, codes_lame):
    """Équivalent Python des formules du classeur (clés = noms de colonnes normalisés)."""
    sortie = {}
    article = str(ligne.get("n° article meule") or "")
    meule = meules.get(article, {})
    sortie["fournisseur"], sortie["ø meule"] = meule.get("fournisseur"), meule.get("diametre")
    sortie["code lame"] = codes_lame.get(norm(str(ligne.get("format") or "") + str(ligne.get("matière") or "")))
    debut, fin = ligne.get("compteur début"), ligne.get("compteur fin")
    haut_i, haut_f = ligne.get("hauteur meule initiale, mm"), ligne.get("hauteur meule finale, mm")
    brut = (fin - debut) if isinstance(fin, (int, float)) and isinstance(debut, (int, float)) and fin != 0 else None
    usure = (haut_i - haut_f) if isinstance(haut_i, (int, float)) and isinstance(haut_f, (int, float)) and haut_i != 0 else None
    sortie["nombre de face brut"] = brut
    sortie["ht.usée, mm"] = usure
    sortie["meule démontée avant usure totale"] = "X" if isinstance(haut_f, (int, float)) and haut_f > 2 else None
    faces_mm = brut / usure if brut is not None and usure else None
    sortie["nb faces/mm"] = faces_mm
    sortie["nb faces/meule de 89 mm"] = faces_mm * 89 if faces_mm is not None else None
    date = ligne.get("date montage")
    sortie["mois"], sortie["année"] = (date.month, date.year) if hasattr(date, "month") else (None, None)
    sortie["ht usée par face, mm"] = usure / brut if brut and usure is not None else None
    return sortie


class ClasseurOpenpyxl:
    def __init__(self, pause_s=0.0):
        self.pause_s = pause_s

    def ouvrir(self, chemin, etape):
        with etape("lancement_excel"):
            time.sleep(self.pause_s)
        with etape("ouverture"):
            try:
                classeur = load_workbook(chemin)
            except PermissionError:
                raise FichierOccupe("Le classeur est utilisé par quelqu'un d'autre : "
                                    "rien n'a été enregistré. Réessayez dans un instant.")
            time.sleep(self.pause_s)
        return ClasseurOpenpyxlOuvert(classeur, chemin, self.pause_s)


class ClasseurOpenpyxlOuvert:
    def __init__(self, classeur, chemin, pause_s):
        self.classeur, self.chemin, self.pause_s = classeur, chemin, pause_s
        self.meules = {}
        for ligne in classeur["Liste Meule"].iter_rows(min_row=2, values_only=True):
            if ligne[0]:
                self.meules[str(ligne[0])] = {"fournisseur": ligne[3], "diametre": ligne[4]}
        self.codes_lame = {}
        for ligne in classeur["Formats & matière"].iter_rows(min_row=2, values_only=True):
            if ligne[2] is not None:
                self.codes_lame[norm(str(ligne[2]) + str(ligne[3] or ""))] = ligne[1]

    def onglets(self):
        return [feuille.title for feuille in self.classeur.worksheets]

    def tableau_data(self):
        if FEUILLE not in self.classeur.sheetnames or TABLEAU not in self.classeur[FEUILLE].tables:
            return None
        return self.classeur[FEUILLE]

    def tableau_de(self, onglet):
        feuille = self.classeur[onglet]
        return feuille if feuille.tables else None

    def colonnes(self, feuille):
        return [cellule.value for cellule in feuille[1]]

    def ajouter(self, feuille, cibles, etape):
        """Ajoute une ligne à la fin de la feuille, agrandit son tableau et renvoie (n° de ligne, colonnes calculées).
        Comme avec Excel, les colonnes calculées sont renvoyées sous leur nom d'en-tête."""
        noms = self.colonnes(feuille)
        entetes = [norm(nom) for nom in noms]
        brut = {norm(nom): valeur for nom, valeur in cibles.items()}
        with etape("ajout_ligne"):
            numero = feuille.max_row + 1
            for table in feuille.tables.values():
                debut, fin = table.ref.split(":")
                fin_colonne = "".join(caractere for caractere in fin if caractere.isalpha())
                table.ref = f"{debut}:{fin_colonne}{numero}"
            time.sleep(self.pause_s)
        with etape("formules"):
            ligne = dict(brut)
            ligne.update({entete: valeur for entete, valeur in _calculer(ligne, self.meules, self.codes_lame).items()
                          if entete in entetes})
            time.sleep(self.pause_s)
        with etape("formats"):
            time.sleep(self.pause_s)
        with etape("valeurs"):
            for colonne, entete in enumerate(entetes, start=1):
                if ligne.get(entete) is not None:
                    feuille.cell(row=numero, column=colonne, value=ligne[entete])
            time.sleep(self.pause_s)
        with etape("calcul"):
            time.sleep(self.pause_s)
        calcules = {str(nom).replace("\n", " "): "" if ligne[entete] is None else ligne[entete]
                    for nom, entete in zip(noms, entetes) if entete not in brut and entete in ligne}
        return numero, calcules

    def enregistrer(self):
        self.classeur.save(self.chemin)

    def fermer(self):
        self.classeur.close()
        time.sleep(self.pause_s)
