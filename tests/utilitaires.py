"""Aides pour les tests : base SQLite temporaire et petit classeur Excel de démonstration."""
import datetime
import os
import shutil
import tempfile
import unittest

from openpyxl import Workbook

from backend import config
from backend.services import donnees, listes, reglages

ENTETES_DATA = ["Matricule monteur", "Format", "Matière", "Face", "Code lame", "N° article meule",
                "Lot fournisseur N° FA / Beleg", "Date montage", "N° de Machine", "MA5", "Pression d'arrosage",
                "Vitesse de rotation", "V.Avance", "Hauteur meule initiale, mm", "Compteur début",
                "Matricule démonteur", "Compteur fin", "Hauteur meule finale, mm", "Commentaire",
                "Nb faces/meule de 89 mm"]


def ligne_data(**modifs):
    """Une ligne de la feuille Data (valeurs plausibles) ; `modifs` remplace des colonnes par leur nom."""
    base = {"Matricule monteur": "0123", "Format": "8", "Matière": "I", "Face": "A", "Code lame": "09999",
            "N° article meule": "0105", "Lot fournisseur N° FA / Beleg": "123", "Date montage": datetime.datetime(2026, 5, 4),
            "N° de Machine": "s7", "MA5": "Non", "Pression d'arrosage": 6, "Vitesse de rotation": 32, "V.Avance": 850,
            "Hauteur meule initiale, mm": 89, "Compteur début": 0, "Matricule démonteur": "0123", "Compteur fin": 1500,
            "Hauteur meule finale, mm": 0, "Commentaire": None, "Nb faces/meule de 89 mm": 1500}
    base.update(modifs)
    return [base[entete] for entete in ENTETES_DATA]


def creer_classeur(chemin, lignes):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(ENTETES_DATA)
    for ligne in lignes:
        ws.append(ligne)
    feuille_formats = wb.create_sheet("Formats & matière")
    feuille_formats.append(["Format", "Matière", "Code article", "MA5", "Meule", "Machine"])
    feuille_formats.append(["8", "I", "09999", "Non", "0105", "s1;s7"])
    feuille_formats.append(["8", "C", "02687", "Peut-être", "0105", "s1"])
    feuille_meules = wb.create_sheet("Liste Meule")
    feuille_meules.append(["Code article", "Etat", "Fournisseur", "Diamètre meule"])
    feuille_meules.append(["0105", "OK", "FA", 300])
    feuille_meules.append(["0999", "OK", "Beleg", 300])
    wb.save(chemin)


class BaseTest(unittest.TestCase):
    """Chaque test travaille dans un dossier temporaire : jamais sur les vraies données."""

    lignes = []          # à remplacer dans les classes qui ont besoin d'un historique

    def setUp(self):
        self.dossier = tempfile.mkdtemp()
        self._anciens = (config.DB_PATH, config.FICHIER)
        config.DB_PATH = os.path.join(self.dossier, "test.sqlite3")
        config.FICHIER = os.path.join(self.dossier, "classeur.xlsx")
        creer_classeur(config.FICHIER, self.lignes)
        for cache in (donnees._cache, donnees._cache_t, listes._cache):
            cache["mtime"] = None
        reglages._cache.update(cle=None, config=None)

    def tearDown(self):
        config.DB_PATH, config.FICHIER = self._anciens
        shutil.rmtree(self.dossier, ignore_errors=True)
