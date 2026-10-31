"""Mode démonstration : classeur fictif et écriture sans Excel."""
import os
import shutil
import tempfile
import unittest

from openpyxl import load_workbook

from backend import config
from backend.services import donnees, listes, reglages
from backend.services.classeur_openpyxl import ClasseurOpenpyxl
from backend.services.ecriture import ajouter_ligne
from backend.services.regles import traiter
from outils.donnees_fictives import creer_classeur


class TestDemo(unittest.TestCase):
    def setUp(self):
        self.dossier = tempfile.mkdtemp()
        self._anciens = (config.DB_PATH, config.FICHIER)
        config.DB_PATH = os.path.join(self.dossier, "demo.sqlite3")
        config.FICHIER = os.path.join(self.dossier, "demo.xlsx")
        creer_classeur(config.FICHIER)
        for cache in (donnees._cache, donnees._cache_t, listes._cache):
            cache["mtime"] = None
        reglages._cache.update(cle=None, config=None)

    def tearDown(self):
        config.DB_PATH, config.FICHIER = self._anciens
        shutil.rmtree(self.dossier, ignore_errors=True)

    def test_classeur_lisible_avec_anomalies(self):
        self.assertEqual(len(donnees.tableau()["lignes"]), 320)
        self.assertEqual(len(donnees.matricules()), 9)       # 8 monteurs + « 534 » (anomalie voulue)

    def test_ajout_complete_les_colonnes_calculees(self):
        saisie = {"Matricule monteur": "1001", "Format": "2", "Matière": "I", "Face": "A", "N° article meule": "0101",
                  "Lot fournisseur N° FA / Beleg": "123", "Date montage": "2026-10-20", "N° de Machine": "s1", "MA5": "Oui",
                  "Pression d'arrosage": 6, "Vitesse de rotation": 32, "V.Avance": 850, "Hauteur meule initiale, mm": 89,
                  "Compteur début": 0, "Matricule démonteur": "1002", "Compteur fin": 2500, "Hauteur meule finale, mm": 0}
        propres, erreurs, _ = traiter(saisie)
        self.assertEqual(erreurs, [])
        etapes = []
        ligne, calcules, info_defaut = ajouter_ligne(propres, defaut=False, progression=etapes.append, mesures={}, classeur=ClasseurOpenpyxl())
        self.assertEqual(ligne, 322)
        self.assertIsNone(info_defaut)
        self.assertIn("sauvegarde", etapes)
        for cache in (donnees._cache, donnees._cache_t):
            cache["mtime"] = None
        derniere = donnees.tableau()["lignes"][-1]
        self.assertEqual(derniere[0], 322)
        self.assertIn("1001", derniere)
        self.assertEqual(load_workbook(config.FICHIER)["Data"].tables["TableauData"].ref[-3:], "322")


if __name__ == "__main__":
    unittest.main()
