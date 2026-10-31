"""Le déroulé de l'écriture d'une saisie dans le classeur, testé par son interface (ajouter_ligne)
avec l'adapter openpyxl : pas besoin d'Excel."""
import glob
import os
import shutil
import tempfile
import unittest
from unittest import mock

from openpyxl import load_workbook

from backend import config
from backend.services import donnees, ecriture, listes, reglages
from backend.services.classeur_openpyxl import ClasseurOpenpyxl
from backend.services.ecriture import FichierOccupe, ajouter_ligne
from backend.services.regles import traiter
from outils.donnees_fictives import creer_classeur

ETAPES = ["attente", "lancement_excel", "ouverture", "ajout_ligne", "formules", "formats", "valeurs", "calcul",
          "sauvegarde", "fermeture"]


def _lignes(feuille):
    """Les lignes d'un onglet du classeur, en dict {en-tête: valeur}."""
    onglet = load_workbook(config.FICHIER)[feuille]
    entetes = [cellule.value for cellule in onglet[1]]
    return [dict(zip(entetes, ligne)) for ligne in onglet.iter_rows(min_row=2, values_only=True)]


class _ClasseurFormuleEnErreur:
    """Faux adapter : délègue tout à l'adapter openpyxl, mais chaque ligne ajoutée a une formule en erreur."""
    def __init__(self):
        self.reel = ClasseurOpenpyxl()

    def ouvrir(self, chemin, etape):
        self.ouvert = self.reel.ouvrir(chemin, etape)
        return self

    def ajouter(self, tableau, cibles, etape):
        numero, _ = self.ouvert.ajouter(tableau, cibles, etape)
        return numero, {"Nb faces/mm": "#DIV/0!"}

    def __getattr__(self, nom):                      # onglets, tableau_data, colonnes, enregistrer, fermer...
        return getattr(self.ouvert, nom)


class TestEcriture(unittest.TestCase):
    def setUp(self):
        self.dossier = tempfile.mkdtemp()
        self.avant = (config.DB_PATH, config.FICHIER, config.DOSSIER_SAUVEGARDES)
        config.DB_PATH = os.path.join(self.dossier, "demo.sqlite3")
        config.FICHIER = os.path.join(self.dossier, "demo.xlsx")
        config.DOSSIER_SAUVEGARDES = os.path.join(self.dossier, "sauvegardes")
        creer_classeur(config.FICHIER)
        for cache in (donnees._cache, donnees._cache_t, listes._cache):
            cache["mtime"] = None
        reglages._cache.update(cle=None, config=None)
        self.classeur = ClasseurOpenpyxl()
        self.saisie = {"Matricule monteur": "1001", "Format": "2", "Matière": "I", "Face": "A", "N° article meule": "0101",
                       "Lot fournisseur N° FA / Beleg": "123", "Date montage": "2026-10-20", "N° de Machine": "s1", "MA5": "Oui",
                       "Pression d'arrosage": 6, "Vitesse de rotation": 32, "V.Avance": 850, "Hauteur meule initiale, mm": 89,
                       "Compteur début": 0, "Matricule démonteur": "1002", "Compteur fin": 2500, "Hauteur meule finale, mm": 0}
        self.defectueuse = {**self.saisie, "Hauteur meule finale, mm": 10, "Commentaire": "Meule cassée ou fissurée : fissure"}

    def tearDown(self):
        config.DB_PATH, config.FICHIER, config.DOSSIER_SAUVEGARDES = self.avant
        shutil.rmtree(self.dossier, ignore_errors=True)

    def propres(self, saisie):
        propres, erreurs, _ = traiter(saisie)
        self.assertEqual(erreurs, [])
        return propres

    def contenu(self):
        with open(config.FICHIER, "rb") as fichier:
            return fichier.read()

    def test_ligne_ajoutee_dans_data_avec_colonnes_calculees(self):
        nb_data = len(_lignes("Data"))
        ligne, calcules, info_defaut = ajouter_ligne(self.propres(self.saisie), classeur=self.classeur)
        self.assertEqual(ligne, nb_data + 2)
        self.assertIsNone(info_defaut)
        self.assertEqual(calcules["Nombre de face brut"], 2500)       # nom d'en-tête, comme avec Excel
        self.assertEqual(_lignes("Data")[-1]["Matricule monteur"], "1001")

    def test_une_copie_de_securite_est_faite_avant_ecriture(self):
        avant = self.contenu()
        ajouter_ligne(self.propres(self.saisie), classeur=self.classeur)
        copies = glob.glob(os.path.join(config.DOSSIER_SAUVEGARDES, "*.xlsx"))
        self.assertEqual(len(copies), 1)
        with open(copies[0], "rb") as fichier:
            self.assertEqual(fichier.read(), avant)

    def test_etapes_annoncees_dans_l_ordre_et_chronometrees(self):
        etapes, mesures = [], {}
        ajouter_ligne(self.propres(self.saisie), progression=etapes.append, mesures=mesures, classeur=self.classeur)
        self.assertEqual(etapes, ETAPES)
        self.assertTrue(set(ETAPES[1:]) | {"total_excel"} <= set(mesures))

    def test_nombre_en_texte_converti_mais_matricule_garde_en_texte(self):
        propres = self.propres(self.saisie)
        propres["Pression d'arrosage"] = "7"
        ajouter_ligne(propres, classeur=self.classeur)
        derniere = _lignes("Data")[-1]
        self.assertEqual(derniere["Pression d'arrosage"], 7)
        self.assertEqual(derniere["Matricule monteur"], "1001")

    def test_colonne_inconnue_refusee_sans_ecriture_partielle(self):
        avant = self.contenu()
        propres = self.propres(self.saisie)
        propres["Colonne qui n'existe pas"] = 1
        with self.assertRaises(ValueError):
            ajouter_ligne(propres, classeur=self.classeur)
        self.assertEqual(self.contenu(), avant)

    def test_meule_defectueuse_dans_data_et_data_defaut(self):
        nb_data, nb_defaut = len(_lignes("Data")), len(_lignes("Data defaut"))
        ligne, _, info = ajouter_ligne(self.propres(self.defectueuse), defaut=True, classeur=self.classeur)
        self.assertEqual(ligne, nb_data + 2)
        self.assertEqual(info["onglet"], "Data defaut")
        self.assertEqual(info["ligne"], nb_defaut + 2)
        self.assertEqual(len(_lignes("Data")), nb_data + 1)
        self.assertEqual(_lignes("Data defaut")[-1]["Commentaire"], "Meule cassée ou fissurée : fissure")

    def test_alias_de_colonne_dans_data_defaut(self):
        classeur_xlsx = load_workbook(config.FICHIER)
        for cellule in classeur_xlsx["Data defaut"][1]:
            if cellule.value == "Matricule monteur":
                cellule.value = "Matricule"                  # alias prévu dans champs.json
        classeur_xlsx.save(config.FICHIER)
        _, _, info = ajouter_ligne(self.propres(self.defectueuse), defaut=True, classeur=self.classeur)
        self.assertNotIn("Matricule monteur", info["colonnes_absentes"])
        self.assertEqual(_lignes("Data defaut")[-1]["Matricule"], "1001")

    def test_meule_defectueuse_seulement_dans_data_defaut(self):
        nb_data, nb_defaut = len(_lignes("Data")), len(_lignes("Data defaut"))
        with mock.patch.dict(ecriture.CONFIG["defaut"], garder_dans_data=False):
            ligne, calcules, info = ajouter_ligne(self.propres(self.defectueuse), defaut=True, classeur=self.classeur)
        self.assertEqual(len(_lignes("Data")), nb_data)
        self.assertEqual(ligne, nb_defaut + 2)
        self.assertEqual(ligne, info["ligne"])
        self.assertEqual(calcules["Nombre de face brut"], 2500)       # colonnes calculées de la ligne Data défaut

    def test_onglet_data_defaut_introuvable_rien_n_est_ecrit(self):
        avant = self.contenu()
        with mock.patch.dict(ecriture.CONFIG["defaut"], onglet="Onglet absent"):
            with self.assertRaises(ValueError):
                ajouter_ligne(self.propres(self.defectueuse), defaut=True, classeur=self.classeur)
        self.assertEqual(self.contenu(), avant)

    def test_onglet_data_defaut_trouve_sans_accent_ni_majuscule(self):
        with mock.patch.dict(ecriture.CONFIG["defaut"], onglet="DATA DÉFAUT"):
            _, _, info = ajouter_ligne(self.propres(self.defectueuse), defaut=True, classeur=self.classeur)
        self.assertEqual(info["onglet"], "Data defaut")

    def test_formule_en_erreur_annule_l_ecriture(self):
        avant = self.contenu()
        with self.assertRaises(ValueError):
            ajouter_ligne(self.propres(self.saisie), classeur=_ClasseurFormuleEnErreur())
        self.assertEqual(self.contenu(), avant)

    def test_tableau_data_absent_rien_n_est_ecrit(self):
        classeur_xlsx = load_workbook(config.FICHIER)
        del classeur_xlsx["Data"].tables["TableauData"]
        classeur_xlsx.save(config.FICHIER)
        avant = self.contenu()
        with self.assertRaises(ValueError):
            ajouter_ligne(self.propres(self.saisie), classeur=self.classeur)
        self.assertEqual(self.contenu(), avant)

    def test_mode_demo_choisit_l_adapter_openpyxl(self):
        with mock.patch.object(config, "MODE_DEMO", True), mock.patch.object(ecriture, "_CLASSEUR", None), \
                mock.patch.object(ecriture, "PAUSE_DEMO_S", 0):
            ligne, _, _ = ajouter_ligne(self.propres(self.saisie))
            self.assertIsInstance(ecriture._CLASSEUR, ClasseurOpenpyxl)
        self.assertEqual(_lignes("Data")[-1]["Matricule monteur"], "1001")
        self.assertEqual(ligne, len(_lignes("Data")) + 1)

    def test_classeur_occupe_libere_le_verrou(self):
        class Occupe:
            def ouvrir(self, chemin, etape):
                raise FichierOccupe("occupé")
        with self.assertRaises(FichierOccupe):
            ajouter_ligne(self.propres(self.saisie), classeur=Occupe())
        self.assertFalse(ecriture.VERROU.locked())
        ajouter_ligne(self.propres(self.saisie), classeur=self.classeur)


class TestAdapterExcel(unittest.TestCase):
    def test_s_importe_sans_excel(self):
        from backend.services import classeur_excel           # pythoncom n'est importé qu'à l'ouverture
        self.assertTrue(hasattr(classeur_excel.ClasseurExcel(), "ouvrir"))


if __name__ == "__main__":
    unittest.main()
