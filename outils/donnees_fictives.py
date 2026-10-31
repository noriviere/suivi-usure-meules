"""Fabrique un classeur Excel de FAUSSES données, avec la même structure que le vrai fichier
(feuilles « Data », « Formats & matière », « Liste Meule », « Data defaut » et leurs tableaux).
Tout est inventé : matricules, fournisseurs, références de meules, valeurs. Résultat reproductible (graine fixe).

    python -m outils.donnees_fictives [chemin.xlsx]
"""
import datetime
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

ENTETES_DATA = [
    "Matricule monteur", "N° article meule", "Fournisseur", "Ø meule", "Lot fournisseur\nN° FA / Beleg", "Date montage",
    "N° de Machine", "MA5", "Code lame", "Format", "Matière", "Face", "Pression d'arrosage", "Vitesse de rotation",
    "V.Avance", "Compteur début", "Matricule démonteur", "Compteur fin", "Hauteur meule initiale, mm",
    "Hauteur meule finale, mm", "Commentaire", "Nombre de face brut", "Ht.Usée, mm", "Meule démontée avant usure totale",
    "Nb faces/mm", "Nb faces/meule de 89 mm", "Mois", "Année", "Ht usée par face, mm"]
ENTETES_DEFAUT = ["Etat", "Remboursement"] + ENTETES_DATA
ENTETES_FORMATS = ["Format+Matière", "Code article", "Format", "Matière", "MA5", "Meule", "Machine"]
ENTETES_MEULES = ["Code article", "Etat", "Commentaire", "Fournisseur", "Diamètre meule", "Epaisseur du bandage meulant",
                  "Epaisseur du bandage polissant", "Hauteur", "Angle", "Prix"]

FOURNISSEURS = ["ABRASIFS DU NORD", "MEULEX", "DIAMANTA"]
MEULES = [  # code, fournisseur, diamètre, prix
    ("0101", 0, 260, 120.0), ("0102", 0, 350, 185.5), ("0103", 1, 350, 190.0), ("0104", 1, 450, 240.0),
    ("0105", 2, 260, 130.0), ("0106", 2, 450, 255.0), ("0107", 0, 450, 235.0), ("0108", 1, 260, 118.0)]
FORMATS = ["2", "4", "6", "8", "10"]
MATIERES = ["I", "C"]
MATRICULES = ["1001", "1002", "1003", "1004", "1005", "1006", "1007", "1008"]
COMMENTAIRES_NON_FINIE = ["Fissure détectée, démontée par précaution", "Meule cassée ou fissurée", "Bruit anormal, arrêt demandé",
                          "Fin de campagne, remplacée avant usure"]


def _referentiel(hasard):
    """Une ligne par couple format + matière : code article, MA5, meule prévue, machines autorisées."""
    lignes = []
    for indice_format, fmt in enumerate(FORMATS):
        for indice_matiere, mat in enumerate(MATIERES):
            rang = indice_format + indice_matiere
            meule = MEULES[(indice_format * 2 + indice_matiere) % len(MEULES)][0]
            machines = ["s%d" % (1 + rang % 6), "s%d" % (1 + (rang + 3) % 6)] if rang % 2 else ["s%d" % (1 + rang % 6)]
            ma5 = "Peut-être" if (indice_format, indice_matiere) == (1, 1) else ("Oui" if rang % 3 == 0 else "Non")
            lignes.append({"format": fmt, "matiere": mat, "code": "%05d" % (1000 + indice_format * 37 + indice_matiere * 11),
                           "ma5": ma5, "meule": meule, "machines": machines})
    return lignes


def _ligne_data(hasard, ref, jour, compteurs, anormale=None):
    code, indice_four, diametre, _prix = next(meule for meule in MEULES if meule[0] == ref["meule"])
    machine = hasard.choice(ref["machines"])
    hauteur_initiale = hasard.choice([89, 89, 89, 86, 84, 80])
    non_finie = anormale == "non_finie" or (anormale is None and hasard.random() < 0.05)
    hauteur_finale = hasard.randint(3, 12) if non_finie else 0
    faces_par_89 = hasard.gauss(1450, 130)
    if anormale == "tres_court":
        faces_par_89 = hasard.uniform(250, 400)
    elif anormale == "tres_long":
        faces_par_89 = hasard.uniform(3100, 3500)
    usure = hauteur_initiale - hauteur_finale
    brut = max(60, int(round(faces_par_89 / 89 * usure)))
    debut = compteurs.get(machine, hasard.randint(500, 4000))
    fin = debut + brut
    compteurs[machine] = fin + hasard.randint(0, 40)
    ma5 = ref["ma5"] if ref["ma5"] != "Peut-être" else hasard.choice(["Oui", "Non"])
    monteur = hasard.choice(MATRICULES)
    demonteur = hasard.choice(MATRICULES)
    if anormale == "matricule_court":
        monteur = "534"
    faces_mm = brut / usure
    return {
        "Matricule monteur": monteur, "N° article meule": ref["meule"], "Fournisseur": FOURNISSEURS[indice_four], "Ø meule": diametre,
        "Lot fournisseur\nN° FA / Beleg": hasard.randint(230000, 239999), "Date montage": datetime.datetime.combine(jour, datetime.time()),
        "N° de Machine": machine, "MA5": ma5, "Code lame": ref["code"], "Format": int(ref["format"]), "Matière": ref["matiere"],
        "Face": hasard.choice(["A", "B"]), "Pression d'arrosage": 0 if hasard.random() < 0.04 else 6,
        "Vitesse de rotation": hasard.randint(30, 36), "V.Avance": hasard.choice([800, 850, 900, 950]),
        "Compteur début": debut, "Matricule démonteur": demonteur, "Compteur fin": fin,
        "Hauteur meule initiale, mm": hauteur_initiale, "Hauteur meule finale, mm": hauteur_finale,
        "Commentaire": hasard.choice(COMMENTAIRES_NON_FINIE) if non_finie else None,
        "Nombre de face brut": brut, "Ht.Usée, mm": usure, "Meule démontée avant usure totale": "X" if hauteur_finale > 2 else None,
        "Nb faces/mm": faces_mm, "Nb faces/meule de 89 mm": faces_mm * 89, "Mois": jour.month, "Année": jour.year,
        "Ht usée par face, mm": usure / brut}


def creer_classeur(chemin, graine=42, nb_lignes=320, aujourdhui=None):
    """Écrit le classeur de démonstration et renvoie son chemin."""
    hasard = random.Random(graine)
    aujourdhui = aujourdhui or datetime.date.today()
    referentiel = _referentiel(hasard)
    wb = Workbook()

    ws = wb.active
    ws.title = "Data"
    ws.append(ENTETES_DATA)
    debut_periode = aujourdhui - datetime.timedelta(days=430)
    anormales = {40: "tres_court", 130: "tres_long", 200: "matricule_court", 260: "tres_court"}
    compteurs = {}
    jours = sorted(debut_periode + datetime.timedelta(days=hasard.randint(0, 425)) for _ in range(nb_lignes))
    for numero, jour in enumerate(jours):
        ligne = _ligne_data(hasard, hasard.choice(referentiel), jour, compteurs, anormales.get(numero))
        ws.append([ligne[entete] for entete in ENTETES_DATA])
    ws.add_table(Table(displayName="TableauData", ref="A1:%s%d" % ("AC", nb_lignes + 1)))

    wf = wb.create_sheet("Formats & matière")
    wf.append(ENTETES_FORMATS)
    for ref in referentiel:
        wf.append([ref["format"] + ref["matiere"], ref["code"], int(ref["format"]), ref["matiere"], ref["ma5"], ref["meule"],
                   ";".join(ref["machines"])])
    wf.add_table(Table(displayName="TableauFormat", ref="A1:G%d" % (len(referentiel) + 1)))

    wm = wb.create_sheet("Liste Meule")
    wm.append(ENTETES_MEULES)
    for code, indice_four, diametre, prix in MEULES:
        wm.append([code, "Prod", None, FOURNISSEURS[indice_four], diametre, 22, 3, 125, 40, prix])
    wm.add_table(Table(displayName="TableauMeules", ref="A1:J%d" % (len(MEULES) + 1)))

    wd = wb.create_sheet("Data defaut")
    wd.append(ENTETES_DEFAUT)
    for _ in range(6):
        jour = aujourdhui - datetime.timedelta(days=hasard.randint(10, 300))
        ligne = _ligne_data(hasard, hasard.choice(referentiel), jour, {}, "non_finie")
        ligne["Commentaire"] = "Meule cassée ou fissurée"
        wd.append(["Défectueuse", None] + [ligne[entete] for entete in ENTETES_DATA])
    wd.add_table(Table(displayName="TableauData3", ref="A1:AE7"))

    os.makedirs(os.path.dirname(os.path.abspath(chemin)), exist_ok=True)
    wb.save(chemin)
    return chemin


if __name__ == "__main__":
    cible = sys.argv[1] if len(sys.argv) > 1 else os.path.join("data", "demo", "classeur_demo.xlsx")
    print("Classeur de démonstration écrit :", creer_classeur(cible))
